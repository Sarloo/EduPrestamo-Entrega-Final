from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Informe-Cierre-EduPrestamo.docx"
ASSETS = ROOT / "docs" / "assets"

INK = "1F2D2B"
MUTED = "586966"
TEAL = "0B675B"
TEAL_PALE = "EAF4F1"
BLUE_PALE = "EEF4F8"
GRAY_PALE = "F5F7F7"
BORDER = "D9D9D9"
WHITE = "FFFFFF"


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)


def set_cell_shading(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=110, start=130, bottom=110, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "6")
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), BORDER)


def set_cell_width(cell, width_inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.append(field_begin)
    run._r.append(instr)
    run._r.append(field_end)


def style_table(table, widths=None, header_fill=TEAL):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for row_index, row in enumerate(table.rows):
        prevent_row_split(row)
        for col_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if widths:
                set_cell_width(cell, widths[col_index])
            if row_index == 0:
                set_cell_shading(cell, header_fill)
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor.from_string(WHITE)
                        run.font.size = Pt(9)
            else:
                if row_index % 2 == 0:
                    set_cell_shading(cell, GRAY_PALE)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(1)
                    paragraph.paragraph_format.line_spacing = 1.08
                    for run in paragraph.runs:
                        run.font.size = Pt(8.7)


def add_table(document, headers, rows, widths=None, header_fill=TEAL):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = str(value)
    style_table(table, widths=widths, header_fill=header_fill)
    document.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_bullets(document, items):
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1.05
        paragraph.add_run(item)


def add_numbered(document, items):
    for item in items:
        paragraph = document.add_paragraph(style="List Number")
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1.05
        paragraph.add_run(item)


def add_caption(document, text):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(10)
    run = paragraph.add_run(text)
    run.italic = True
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string(MUTED)


def add_picture(document, filename, caption):
    path = ASSETS / filename
    if not path.exists():
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(path), width=Inches(5.8))
    add_caption(document, caption)


def configure_document(document):
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(1.3)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.35)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.16

    title = styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(29)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(14)
    title_p_pr = title._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    for style_name, size, before, after in (
        ("Heading 1", 17, 16, 8),
        ("Heading 2", 13, 12, 6),
        ("Heading 3", 11, 9, 4),
    ):
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

def add_cover(document):
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(70)

    eyebrow = document.add_paragraph()
    eyebrow.alignment = WD_ALIGN_PARAGRAPH.CENTER
    eyebrow.paragraph_format.space_after = Pt(12)
    run = eyebrow.add_run("PROYECTO ESCOLAR DE INGENIERÍA DE SOFTWARE")
    run.font.name = "Arial"
    run.font.size = Pt(9)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(TEAL)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Informe de cierre EduPréstamo")
    title_p_pr = title._p.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(40)
    run = subtitle.add_run("Sistema para la gestión de recursos y préstamos escolares")
    run.font.name = "Arial"
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor.from_string(MUTED)

    summary = document.add_paragraph()
    summary.alignment = WD_ALIGN_PARAGRAPH.CENTER
    summary.paragraph_format.left_indent = Inches(0.65)
    summary.paragraph_format.right_indent = Inches(0.65)
    summary.paragraph_format.space_after = Pt(42)
    summary.add_run(
        "La entrega implementa autenticación JWT, roles ADMIN y USER, inventario, "
        "solicitudes de préstamo, pruebas automatizadas, CI/CD y análisis de seguridad "
        "y calidad. Las 27 pruebas aprobaron y la cobertura superó el 80 por ciento."
    )

    metadata = document.add_paragraph()
    metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER
    metadata.paragraph_format.space_before = Pt(80)
    metadata.add_run("Fecha de cierre\n").bold = True
    metadata.add_run("18 de septiembre de 2026\n\n")
    metadata.add_run("Versión\n").bold = True
    metadata.add_run("1.0.0")

    document.add_page_break()


def build_report():
    document = Document()
    configure_document(document)
    add_cover(document)

    document.add_heading("Resumen ejecutivo", level=1)
    document.add_paragraph(
        "EduPréstamo quedó implementado como una aplicación web funcional para registrar "
        "recursos escolares y controlar el ciclo de préstamo. El sistema permite crear "
        "cuentas USER, iniciar sesión mediante JWT, consultar disponibilidad, solicitar "
        "materiales y administrar aprobaciones, entregas y devoluciones. El administrador "
        "también dispone de indicadores, gestión de recursos y trazabilidad por solicitante."
    )
    document.add_paragraph(
        "La validación automática aprobó 27 pruebas en cuatro suites. Jest registró 94.86 % "
        "de statements, 85.71 % de branches, 98.64 % de functions y 96.02 % de lines. "
        "OWASP ZAP no encontró XSS, inyección SQL ni alertas altas. SonarQube aprobó el "
        "Quality Gate con cero bugs, vulnerabilidades, code smells, deuda técnica y duplicación."
    )

    document.add_heading("Objetivo del proyecto", level=1)
    document.add_paragraph(
        "Implementar y validar, antes del cierre del proyecto escolar, un módulo web de "
        "préstamos de recursos con autenticación JWT, roles ADMIN y USER, cobertura automatizada "
        "mínima de 80 %, pipeline CI/CD, análisis OWASP ZAP y análisis SonarQube, conservando "
        "evidencias reproducibles de cada resultado."
    )
    document.add_paragraph(
        "El objetivo es específico porque define el módulo y sus controles; medible por la "
        "cobertura, pruebas y Quality Gate; alcanzable mediante una arquitectura modular; "
        "relevante para la administración de recursos; y delimitado por la fecha de cierre."
    )

    document.add_heading("Alcance funcional", level=1)
    add_bullets(document, [
        "Registro de cuentas públicas con rol USER fijo e inicio de sesión con JWT.",
        "Control de acceso para funciones administrativas y funciones personales.",
        "Inventario con categorías, condición, existencias, búsqueda y baja lógica.",
        "Solicitudes con estados PENDING, APPROVED, REJECTED, DELIVERED, RETURNED y CANCELLED.",
        "Actualización transaccional de existencias durante aprobación y devolución.",
        "Panel administrativo con resumen, recursos, solicitudes y datos del solicitante.",
        "Reportes técnicos de pruebas, seguridad, calidad y despliegue.",
    ])

    document.add_page_break()
    document.add_heading("Arquitectura implementada", level=1)
    document.add_paragraph(
        "La solución utiliza una SPA en HTML, CSS y JavaScript, una API REST en Node.js 22 "
        "con Express 5 y una base SQLite. El prototipo se mantuvo como monolito modular para "
        "reducir dependencias externas y facilitar su evaluación. Las consultas SQL usan "
        "parámetros y las operaciones que afectan inventario se ejecutan dentro de transacciones."
    )
    add_table(
        document,
        ["Capa", "Tecnología", "Responsabilidad"],
        [
            ["Interfaz", "HTML CSS JavaScript", "Acceso, catálogo, solicitudes y administración"],
            ["API", "Node.js 22 y Express 5", "Reglas de negocio, validación y respuestas JSON"],
            ["Datos", "SQLite", "Usuarios, recursos, préstamos y trazabilidad"],
            ["Seguridad", "JWT bcrypt Helmet CORS", "Autenticación, roles, contraseñas y encabezados"],
            ["Entrega", "Docker GitHub Actions", "Entorno reproducible y validación automatizada"],
        ],
        widths=[1.0, 1.75, 3.75],
    )
    add_picture(document, "ui-admin-summary.png", "Figura 1 Panel administrativo y resumen operativo")

    document.add_heading("Implementación y seguridad", level=1)
    document.add_heading("Autenticación y roles", level=2)
    document.add_paragraph(
        "El servidor firma tokens JWT con HS256 y valida emisor, audiencia y caducidad. Las "
        "contraseñas se almacenan con bcrypt. El registro público no acepta un rol enviado por "
        "el cliente y asigna siempre USER. Las rutas administrativas exigen el rol ADMIN."
    )
    document.add_heading("Protecciones aplicadas", level=2)
    add_bullets(document, [
        "Bloqueo temporal de la cuenta después de tres intentos fallidos.",
        "Respuesta uniforme de autenticación para evitar enumeración de correos.",
        "Consultas SQL parametrizadas y transacciones para cambios de existencias.",
        "Helmet con política de seguridad de contenido y encabezados adicionales.",
        "CORS del mismo origen por defecto y cuerpo JSON limitado a 32 KiB.",
        "Validación de longitud, formato, fechas, cantidades, estados y roles.",
        "Secreto JWT obligatorio y de al menos 32 bytes en producción.",
    ])
    document.add_heading("Separación de información", level=2)
    document.add_paragraph(
        "Cada usuario consulta únicamente sus préstamos. El administrador conserva dos vistas: "
        "Mis préstamos muestra sus solicitudes personales y Solicitudes muestra todos los "
        "registros con nombre y correo del solicitante. Esta separación evita confundir la "
        "actividad administrativa con la actividad personal."
    )
    add_picture(document, "ui-user-catalog.png", "Figura 2 Catálogo disponible para una cuenta USER")

    document.add_page_break()
    document.add_heading("Pruebas automatizadas", level=1)
    document.add_paragraph(
        "Las pruebas se ejecutaron con Jest y Supertest en una instalación limpia. Cubren "
        "autenticación, autorización, privacidad entre cuentas, recursos, préstamos, estados, "
        "sobreasignación, reportes y manejo de errores. El umbral global configurado es 80 %."
    )
    add_table(
        document,
        ["Métrica", "Resultado", "Umbral", "Estado"],
        [
            ["Statements", "94.86 %", "80 %", "Aprobado"],
            ["Branches", "85.71 %", "80 %", "Aprobado"],
            ["Functions", "98.64 %", "80 %", "Aprobado"],
            ["Lines", "96.02 %", "80 %", "Aprobado"],
        ],
        widths=[1.8, 1.55, 1.55, 1.6],
    )
    document.add_paragraph(
        "Resultado de ejecución: 4 suites aprobadas, 27 pruebas aprobadas y 0 fallidas. "
        "La cobertura publicada corresponde al backend y no incluye src/public, "
        "src/server.js ni src/seed.js."
    )

    document.add_heading("Integración y entrega continuas", level=1)
    document.add_paragraph(
        "El workflow de GitHub Actions se ejecuta en pushes, pull requests y de forma manual. "
        "La primera etapa instala dependencias, ejecuta ESLint y Jest, y publica la cobertura. "
        "La segunda construye la imagen, inicia un entorno efímero, espera el healthcheck y "
        "ejecuta ZAP. La tercera envía el análisis a Sonar cuando existe un token configurado."
    )
    add_numbered(document, [
        "Quality ejecuta npm ci, lint y pruebas con cobertura.",
        "Container security construye y valida la imagen de producción.",
        "El despliegue de prueba responde en GET /health.",
        "OWASP ZAP escanea el entorno autorizado y conserva sus reportes.",
        "Sonar analiza fuentes, pruebas y LCOV cuando se configura SONAR_TOKEN.",
    ])
    document.add_paragraph(
        "Las etapas principales fueron validadas localmente. El repositorio aún no está "
        "conectado a GitHub, por lo que no se declara una ejecución remota ni una URL inexistente. "
        "Al publicar el repositorio, el siguiente push iniciará el workflow configurado."
    )

    document.add_heading("Resultados de OWASP ZAP", level=1)
    add_table(
        document,
        ["Severidad", "Tipos de alerta", "Interpretación"],
        [
            ["Alta", "0", "Sin hallazgos"],
            ["Media", "2", "CSRF sin cookies y HTTP local"],
            ["Baja", "0", "Sin hallazgos"],
            ["Informativa", "4", "Autenticación y comportamiento de caché"],
        ],
        widths=[1.35, 1.35, 3.8],
    )
    document.add_paragraph(
        "ZAP 2.17.0 no detectó alertas de XSS, inyección SQL ni severidad alta. La alerta anti "
        "CSRF se revisó considerando que la autenticación se transmite mediante JWT Bearer y no "
        "mediante una cookie de sesión automática. La alerta HTTP corresponde exclusivamente al "
        "entorno local. Un despliegue público debe usar TLS."
    )
    document.add_paragraph(
        "El spider cubrió la superficie pública y el endpoint de inicio de sesión. Las rutas "
        "protegidas también se ejercitaron mediante pruebas HTTP automatizadas. Se conserva como "
        "mejora pendiente un escaneo API autenticado para USER y ADMIN."
    )

    document.add_heading("Resultados de SonarQube", level=1)
    add_table(
        document,
        ["Métrica", "Primer análisis", "Resultado final"],
        [
            ["Quality Gate", "Aprobado", "Aprobado"],
            ["Bugs", "0", "0"],
            ["Vulnerabilidades", "0", "0"],
            ["Code smells", "7", "0"],
            ["Deuda técnica", "63 minutos", "0 minutos"],
            ["Duplicación", "0.0 %", "0.0 %"],
            ["Cobertura Sonar", "91.7 %", "91.6 %"],
            ["Hotspots", "4", "4 revisados"],
        ],
        widths=[2.35, 2.05, 2.1],
    )
    document.add_paragraph(
        "La refactorización eliminó los siete code smells y toda la deuda técnica reportada. "
        "Sonar mantuvo calificaciones A en fiabilidad, seguridad y mantenibilidad. Los cuatro "
        "hotspots corresponden a dos credenciales demo, un falso positivo en la función de "
        "validación de contraseña y una expresión regular limitada a entradas de 254 caracteres."
    )

    document.add_heading("Comparación entre planificación y ejecución", level=1)
    add_table(
        document,
        ["Elemento", "Planificado", "Ejecutado", "Causa e impacto"],
        [
            ["Base de datos", "MySQL", "SQLite", "Facilitó una demostración reproducible sin servidor externo. La migración queda prevista."],
            ["Roles", "Estudiante docente administrador", "USER y ADMIN", "Se consolidaron perfiles operativos para cumplir el alcance básico sin duplicar permisos."],
            ["Módulos", "Inventario préstamos historial reportes", "Implementados", "El historial se representa mediante estados y marcas de tiempo de cada préstamo."],
            ["Cobertura", "Mínimo 80 %", "85.71 % a 98.64 % según métrica", "El umbral se superó en todas las métricas de backend."],
            ["Seguridad", "JWT roles ZAP", "Implementado y analizado", "Los hallazgos locales quedaron documentados con su tratamiento."],
            ["Calidad", "SonarQube", "Quality Gate aprobado", "La primera revisión permitió eliminar siete code smells."],
            ["CI CD", "Automatización y despliegue de prueba", "Workflow y entorno efímero implementados", "La ejecución remota requiere conectar el repositorio a GitHub."],
            ["Registro", "Módulo básico de personas", "Autoservicio USER", "El servidor impide que el cliente se asigne privilegios ADMIN."],
        ],
        widths=[1.05, 1.55, 1.55, 2.35],
        header_fill="314F63",
    )

    document.add_heading("Lecciones aprendidas", level=1)
    document.add_heading("La evidencia debe generarse durante el desarrollo", level=2)
    document.add_paragraph(
        "Configurar pruebas, cobertura y reportes desde el inicio permite detectar regresiones y "
        "evita reconstruir resultados al cierre. Los archivos originales son más confiables que "
        "capturas aisladas porque conservan métricas, versiones y detalles técnicos."
    )
    document.add_heading("La seguridad requiere revisión humana", level=2)
    document.add_paragraph(
        "ZAP y Sonar identifican señales, pero su severidad depende del contexto. La alerta CSRF "
        "cambia de tratamiento según el mecanismo de sesión, y los hotspots de credenciales demo "
        "no equivalen automáticamente a secretos de producción. Cada hallazgo necesita una "
        "decisión documentada."
    )
    document.add_heading("Los roles deben reflejarse en la interfaz", level=2)
    document.add_paragraph(
        "Una API puede filtrar datos correctamente y aun así confundir al usuario si la interfaz "
        "usa la ruta equivocada. Separar Mis préstamos de Solicitudes administrativas aclaró la "
        "propiedad de cada registro y mejoró la trazabilidad."
    )
    document.add_heading("La portabilidad influye en la arquitectura", level=2)
    document.add_paragraph(
        "SQLite simplificó la instalación y las pruebas del prototipo. La decisión reduce la "
        "carga operativa de la entrega, aunque un sistema institucional con concurrencia y copias "
        "de seguridad deberá migrar a un servicio de base de datos administrado."
    )

    document.add_heading("Plan de mejora continua", level=1)
    add_table(
        document,
        ["Plazo", "Mejora", "Acción", "Indicador"],
        [
            ["30 días", "GitHub y protección de ramas", "Publicar el repositorio, configurar secretos y exigir el workflow en cada pull request.", "100 % de cambios con pipeline aprobado"],
            ["30 días", "Sesiones", "Migrar el JWT de localStorage a cookie HttpOnly Secure SameSite y agregar rate limiting.", "0 tokens accesibles desde JavaScript y respuesta 429 probada"],
            ["60 días", "Seguridad API", "Publicar OpenAPI y ejecutar ZAP autenticado para USER y ADMIN.", "0 alertas altas y 0 medias sin revisión"],
            ["90 días", "Persistencia", "Migrar a MySQL o PostgreSQL administrado y automatizar respaldos.", "RPO máximo de 24 horas y restauración probada"],
            ["90 días", "Interfaz", "Agregar pruebas de navegador para registro, roles y flujo completo de préstamo.", "Flujos críticos aprobados en cada release"],
            ["6 meses", "Predicción de demanda", "Evaluar un modelo que anticipe recursos con mayor demanda usando historial suficiente.", "MAPE menor o igual a 20 % frente a una línea base"],
        ],
        widths=[0.85, 1.35, 2.9, 1.4],
        header_fill="314F63",
    )

    document.add_heading("Limitaciones y riesgos restantes", level=1)
    add_bullets(document, [
        "El pipeline remoto todavía no tiene una ejecución porque falta conectar el repositorio a GitHub.",
        "El escaneo ZAP público no reemplaza un escaneo autenticado de todos los endpoints protegidos.",
        "El JWT se conserva en localStorage y debe migrarse antes de un despliegue público.",
        "Las credenciales demo deben eliminarse fuera de la presentación académica.",
        "SQLite es adecuado para el prototipo, pero requiere una estrategia adicional de concurrencia y respaldo en producción.",
    ])

    document.add_heading("Cierre", level=1)
    document.add_paragraph(
        "EduPréstamo cumple el alcance funcional y técnico solicitado para la entrega académica. "
        "La aplicación implementa autenticación JWT, roles, registro USER, inventario y préstamo "
        "de recursos. Las pruebas superan el umbral de cobertura y las herramientas de seguridad "
        "y calidad producen evidencia verificable. Las limitaciones pendientes están convertidas "
        "en acciones medibles dentro del plan de mejora."
    )

    document.add_page_break()
    document.add_heading("Ubicación de las evidencias", level=1)
    add_table(
        document,
        ["Evidencia", "Ruta dentro del repositorio"],
        [
            ["Cobertura Jest", "reports/coverage/"],
            ["Resumen de pruebas", "reports/unit/test-summary.md"],
            ["OWASP ZAP", "reports/security/zap/"],
            ["SonarQube", "reports/quality/sonarqube/"],
            ["CI CD", "reports/cicd/local-validation.md"],
            ["Despliegue", "reports/deployment/container-evidence.md"],
            ["Workflow", ".github/workflows/ci-cd.yml"],
        ],
        widths=[2.1, 4.4],
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_report()
